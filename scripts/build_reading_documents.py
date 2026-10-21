"""Synchronise Word reading copies with the continuous public edition."""
from pathlib import Path
import sys
import json
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt,RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT
R=Path(__file__).resolve().parents[1]
h=json.loads((R/'hyperbook-data.json').read_text());records={r['id']:r for r in h['records']}
def source_link(p,text,url):
 link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
 run=OxmlElement('w:r');properties=OxmlElement('w:rPr')
 for tag,attrs in [('w:rFonts',{'w:ascii':'Courier New','w:hAnsi':'Courier New'}),('w:sz',{'w:val':'17'}),('w:color',{'w:val':'000000'})]:
  element=OxmlElement(tag)
  for key,value in attrs.items():element.set(qn(key),value)
  properties.append(element)
 run.append(properties);value=OxmlElement('w:t');value.text=text;run.append(value);link.append(run);p._p.append(link)
for g in h['groups']:
 if len(sys.argv)>1 and g['number']!=int(sys.argv[1]):continue
 path=R/g['path']/'reading.docx';doc=Document(path)
 for child in list(doc._element.body):
  if child.tag!=qn('w:sectPr'):doc._element.body.remove(child)
 doc.core_properties.title='Summa Transscriptus · '+g['guest'];doc.core_properties.author='Organ for de Autonome Videnskaber';doc.core_properties.last_modified_by='Organ for de Autonome Videnskaber'
 for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
  st=doc.styles[name];st.font.color.rgb=RGBColor(0,0,0)
 doc.add_paragraph('Summa Transscriptus','Title');doc.add_paragraph(g['guest']+' og Computer Lars','Subtitle');doc.add_paragraph(g['date']+' · Under Bugten · Aarhus Bugt Center for Kunst på Det Jyske Kunstakademi')
 if h.get('performance',{}).get('closing_signal'):
  p=doc.add_paragraph(h['performance']['closing_signal']);p.paragraph_format.space_after=Pt(16)
  for run in p.runs:run.font.size=Pt(9);run.italic=True
 for rid in g['records']:
  r=records[rid];title=r['title'] if r['title']==r['position'] or r['title'].startswith(r['position']+' ') else r['position']+' · '+r['title'];doc.add_paragraph(title,'Heading 1');p=doc.add_paragraph(rid+' · '+g['date']);p.paragraph_format.space_after=Pt(18)
  for t in r['utterances']:
   sec=int(t['start']);clock=f'{sec//3600:02}:{sec%3600//60:02}:{sec%60:02}'
   p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(12);run=p.add_run(t['speaker']+' · '+clock);run.font.size=Pt(9);run.italic=True
   for b in t['blocks']:
    doc.add_paragraph(b)
    if rid=='D1-01' and b.startswith('Incipere reconstructionem'):
     p=doc.add_paragraph('Oratio latina · rekonstrueret forlæg efter Hans-Jørgen Schanz, Til rekonstruktionen af kritikken af den politiske økonomis omfangslogiske status. Aarhus: Modtryk, 1973.')
     p.paragraph_format.space_after=Pt(12)
     for run in p.runs:run.font.size=Pt(8.5);run.italic=True
  if h['record_references'].get(rid):
   p=doc.add_paragraph('Henvisninger');p.paragraph_format.keep_with_next=True
   p.paragraph_format.space_before=Pt(12)
   for run in p.runs:run.font.size=Pt(9);run.italic=True
   for ref in h['record_references'][rid]:
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(6)
    url=ref['url']if ref['url'].startswith('http')else 'https://computerlars.github.io/Summa-Transscriptus/'+ref['url']
    source_link(p,'['+str(ref['number'])+'] '+ref['display'],url)
  history=h['collation']['records'][rid].get('historical')
  if history:
   doc.add_paragraph('Maskinsyntese · 2023','Heading 2');doc.add_paragraph(history['synthesis'])
  elif rid=='D4-21':
   doc.add_paragraph('Maskinsyntese · 2023','Heading 2');doc.add_paragraph('Ingen maskinsyntese er bevaret for Metatese 21.')
  doc.paragraphs[-1].paragraph_format.keep_with_next=True
  p=doc.add_paragraph('Læseudgave og lydsteder: https://computerlars.github.io/Summa-Transscriptus/#disput/'+rid);p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(6)
  for run in p.runs:run.font.size=Pt(8.5)
 doc.save(path)
 print(path.relative_to(R),len(doc.paragraphs),'paragraphs')
